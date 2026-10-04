import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

# Colors
NAVY = RGBColor(10, 17, 32)        # #0A1120
TEAL = RGBColor(13, 148, 136)      # #0D9488
SKY = RGBColor(2, 132, 199)        # #0284C7
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
SLATE = RGBColor(100, 116, 139)

SLIDES_DATA = [
    {
        "title": "SACH AI",
        "subtitle": "Smart Assistant for Checking Honesty in Financial Information\n\"Suniye, Samajhiye, Sachai pe Bharosa Kijiye.\"\n\nSANGYAN Hackathon • Track E (Misinformation & Content Literacy)\nOrganized by SNTC, IIT (BHU) Varanasi in collaboration with SEBI and NSDL",
        "points": []
    },
    {
        "title": "1. Problem Statement: Retail Investor Vulnerability",
        "subtitle": "Unprecedented surge in retail market participation meets organized cyber deception:",
        "points": [
            "Over 150 Million retail Demat accounts in India with exponential growth across Tier-2/Tier-3 cities.",
            "Proliferation of unregulated 'VIP' Telegram channels, WhatsApp stock tip syndicates, and fake IPO allotment links.",
            "Deceptive AI deepfakes and impersonation videos targeting unsuspecting retail investors with 300% profit promises.",
            "Critical lack of accessible, native-language verification tools tailored for non-technical or elderly citizens."
        ]
    },
    {
        "title": "2. Target Users & Core Challenges",
        "subtitle": "Empowering those most exposed to financial misinformation:",
        "points": [
            "First-Time Retail Investors: Vulnerable to FOMO, false guarantees, and operator penny stock pumps.",
            "Elderly & Senior Citizens: Frequently targeted by high-pressure phishing messages and account freeze threats.",
            "Regional Language Users: Millions who communicate in Hindi, Bengali, or regional dialects with zero access to English regulatory portals.",
            "Voice-Note Consumers: Users who receive financial advice via WhatsApp voice clips rather than formal prospectuses."
        ]
    },
    {
        "title": "3. Proposed Solution: SACH AI",
        "subtitle": "An accessible, voice-first multilingual misinformation & content literacy shield:",
        "points": [
            "Voice-First Ingestion: Record via browser mic, upload audio notes (.wav/.mp3), or paste text directly.",
            "Multilingual Support: Full UI and speech pipeline in English, Hindi (हिन्दी), and Bengali (বাংলা).",
            "Hybrid AI & NLP Rule Engine: Real-time identification of prohibited guaranteed returns and dabba trading.",
            "Spoken Plain-Language Explanations: Translates complex regulatory principles into clear audio guidance.",
            "Evidence-Aware Linking: Directly cross-references official SEBI SCORES, NSDL, BSE/NSE, and 1930 Helpline."
        ]
    },
    {
        "title": "4. System Architecture & Workflow Pipeline",
        "subtitle": "End-to-end processing from spoken audio to verified evidence:",
        "points": [
            "1. Audio & Text Capture: Ingests audio streams with real-time waveform and duration monitoring.",
            "2. Speech-to-Text: Robust transcription engine with language detection and editable transcript review.",
            "3. NLP Claim Analyzer: Tokenizes sentences, extracts specific claims, tags red flag triggers, and calculates 0-100 risk score.",
            "4. Regulatory Evidence Mapping: Cites statutory mechanisms (ASBA for IPOs, SCORES for broker grievances).",
            "5. Text-to-Speech Output: Generates base64 voice responses (gTTS) for auditory accessibility."
        ]
    },
    {
        "title": "5. Technology Stack & Implementation",
        "subtitle": "Modern, robust, modular full-stack architecture:",
        "points": [
            "Frontend: React.js 18 with Vite, Tailwind CSS, Lucide Icons, responsive mobile-first UI.",
            "Backend: Python 3.12, FastAPI asynchronous framework, Pydantic validation, Uvicorn server.",
            "AI & Speech: SpeechRecognition, gTTS, modular NLP rules engine, optional OpenAI/Gemini integration.",
            "Database & Telemetry: SQLite store for non-sensitive aggregated audit logs.",
            "Deployment: Docker & Docker Compose containerization with environment variable secret management."
        ]
    },
    {
        "title": "6. Accessibility, Privacy & Safety by Design",
        "subtitle": "Engineered for public trust and absolute user safety:",
        "points": [
            "Accessibility Modes: Built-in Large Text mode (A+) and High-Contrast mode for elderly and low-vision users.",
            "Zero Persistent Audio Retention: Audio files are processed in-memory / temporary buffers and immediately purged.",
            "Strict Non-Advisory Guardrail: Programmed never to provide stock recommendations, buy/sell targets, or price tips.",
            "No False Guarantees: Always emphasizes statutory market risk even on safe informational statements."
        ]
    },
    {
        "title": "7. Regulatory & Official Evidence Integration",
        "subtitle": "Seamless alignment with Indian investor protection institutions:",
        "points": [
            "SEBI SCORES: Direct links for lodging complaints against fraudulent brokers and unlisted advisors.",
            "SEBI Saa₹thi: Educational guides on mutual funds and securities market fundamentals.",
            "NSDL Depository Services: Demat credit and Consolidated Account Statement (CAS) verification.",
            "National Cyber Crime Portal & 1930 Helpline: 4-step emergency action plan to freeze defrauded funds in transit."
        ]
    },
    {
        "title": "8. Expected Impact & Scalability",
        "subtitle": "Protecting India's next 100 million retail investors:",
        "points": [
            "Democratizing Investor Literacy: Bringing regulatory wisdom to regional language speakers across Bharat.",
            "Proactive Fraud Prevention: Enabling instant 10-second verification before transferring funds to fraudsters.",
            "Extensible Microservice: Designed for plug-and-play integration as a broker API or WhatsApp Bot."
        ]
    },
    {
        "title": "9. Roadmap & Conclusion",
        "subtitle": "The path forward for investor protection technology:",
        "points": [
            "Phase 1 (Completed Prototype): Multilingual web app with voice I/O, scam quiz, and evidence mapping.",
            "Phase 2: Official WhatsApp and Telegram Bot integration for inline forward fact-checking.",
            "Phase 3: Multimodal video deepfake detection and screenshot OCR analyzer.",
            "Tagline: \"Suniye, Samajhiye, Sachai pe Bharosa Kijiye.\""
        ]
    }
]

def generate_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for slide_idx, item in enumerate(SLIDES_DATA):
        slide = prs.slides.add_slide(blank_layout)
        
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        if slide_idx == 0:
            bg.fill.solid()
            bg.fill.fore_color.rgb = NAVY
        else:
            bg.fill.solid()
            bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.color.rgb = NAVY

        # Header / Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item["title"]
        p.font.size = Pt(32 if slide_idx > 0 else 44)
        p.font.bold = True
        p.font.color.rgb = WHITE if slide_idx == 0 else NAVY

        # Subtitle Box
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.0))
        stf = sub_box.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = item["subtitle"]
        sp.font.size = Pt(18 if slide_idx == 0 else 16)
        sp.font.color.rgb = TEAL if slide_idx == 0 else SLATE

        # Points
        if item["points"]:
            points_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(4.0))
            ptf = points_box.text_frame
            ptf.word_wrap = True
            for pt_text in item["points"]:
                pp = ptf.add_paragraph()
                pp.text = f"•  {pt_text}"
                pp.font.size = Pt(17)
                pp.font.color.rgb = NAVY
                pp.space_after = Pt(14)

    prs.save(output_path)
    print(f"Generated PPTX at {output_path}")

def generate_pdf(output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CoverTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=28,
        leading=34,
        textColor=colors.HexColor('#0A1120'),
        alignment=0
    )
    
    sub_style = ParagraphStyle(
        'CoverSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=13,
        leading=18,
        textColor=colors.HexColor('#0D9488'),
        spaceAfter=15
    )
    
    body_style = ParagraphStyle(
        'SlideBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=18,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=8
    )

    story = []
    for idx, item in enumerate(SLIDES_DATA):
        story.append(Paragraph(item["title"], title_style))
        story.append(Spacer(1, 8))
        story.append(Paragraph(item["subtitle"].replace('\n', '<br/>'), sub_style))
        story.append(Spacer(1, 10))
        
        for pt in item["points"]:
            story.append(Paragraph(f"<b>&bull;</b> {pt}", body_style))
        
        if idx < len(SLIDES_DATA) - 1:
            story.append(PageBreak())

    doc.build(story)
    print(f"Generated PDF at {output_path}")

if __name__ == "__main__":
    os.makedirs("presentation", exist_ok=True)
    generate_pptx("presentation/SACH_AI_Pitch_Deck.pptx")
    generate_pdf("presentation/SACH_AI_Pitch_Deck.pdf")
